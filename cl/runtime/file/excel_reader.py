# Copyright (C) 2023-present The Project Contributors
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import csv
import datetime as dt
import os
import re
import types
from dataclasses import dataclass
from typing import Sequence
from typing import Union
from typing import get_args
from typing import get_origin
from typing import get_type_hints
from frozendict import frozendict
from openpyxl import load_workbook
from cl.runtime.file.file_util import FileUtil
from cl.runtime.file.reader import Reader
from cl.runtime.primitive.case_util import CaseUtil
from cl.runtime.records.record_mixin import RecordMixin

_INVALID_SHEET_NAME_RE = re.compile(r'[/\\<>:"|?*\x00\n]')


@dataclass(slots=True, kw_only=True)
class ExcelReader(Reader):
    """Read records from XLSX files by converting to CSV files on disk."""

    def load_all(
        self,
        *,
        dirs: Sequence[str],
        ext: str,
        file_include_patterns: Sequence[str] | None = None,
        file_exclude_patterns: Sequence[str] | None = None,
    ) -> frozendict[str, tuple[RecordMixin, ...]]:
        # Conversion is done separately via convert_to_csv, records are loaded by CsvReader
        return frozendict()

    @classmethod
    def convert_to_csv(
        cls,
        *,
        dirs: Sequence[str],
        ext: str,
        file_include_patterns: Sequence[str] | None = None,
        file_exclude_patterns: Sequence[str] | None = None,
    ) -> None:
        """Convert all XLSX files to CSV files on disk."""

        file_paths = FileUtil.enumerate_files(
            dirs=dirs,
            ext=ext,
            file_include_patterns=file_include_patterns,
            file_exclude_patterns=file_exclude_patterns,
        )

        for xlsx_path in file_paths:
            cls._convert_file(xlsx_path)

    @classmethod
    def _convert_file(cls, xlsx_path: str) -> list[str]:
        """Convert a single XLSX file to one or more CSV files.

        Returns:
            List of generated CSV file paths.
        """

        try:
            return cls._convert_file_impl(xlsx_path)
        except Exception as e:
            raise RuntimeError(f"Failed to convert XLSX file {xlsx_path}.\nError: {e}") from e

    @classmethod
    def _convert_file_impl(cls, xlsx_path: str) -> list[str]:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        sheet_names = wb.sheetnames

        stem = os.path.splitext(xlsx_path)[0]
        dir_path = os.path.dirname(xlsx_path)
        base_name = os.path.splitext(os.path.basename(xlsx_path))[0]

        # Normalize names and check for collisions
        normalized = {}
        for name in sheet_names:
            norm_name = cls._normalize_sheet_name(name)
            if norm_name in normalized:
                raise RuntimeError(
                    f"Sheet name collision in '{xlsx_path}': sheets '{normalized[norm_name]}' and '{name}' "
                    f"both normalize to '{norm_name}'."
                )
            normalized[norm_name] = name

        # Get record type and temporal field types from schema
        record_type = FileUtil.get_type_from_filename(xlsx_path, raise_on_fail=False)
        temporal_fields = cls._get_temporal_field_types(record_type) if record_type else {}

        csv_paths = []
        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # First row is headers
            headers = [str(cell) if cell is not None else "" for cell in rows[0]]
            data_rows = rows[1:]

            # Map column index to temporal field type using PascalCase-to-snake_case conversion
            column_types: dict[int, type] = {}
            for i, header in enumerate(headers):
                if header.startswith("_"):
                    continue
                field_name = CaseUtil.pascal_to_snake_case(header)
                if field_name in temporal_fields:
                    column_types[i] = temporal_fields[field_name]

            # Create CSV filenames
            norm_name = cls._normalize_sheet_name(sheet_name)
            csv_path = os.path.join(dir_path, f"{base_name}.{norm_name}.csv")

            # Write CSV with consistent settings
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(
                    f,
                    delimiter=",",
                    quotechar='"',
                    quoting=csv.QUOTE_MINIMAL,  # noqa Expects a literal
                    lineterminator=os.linesep,
                )
                writer.writerow(headers)
                for row in data_rows:
                    writer.writerow(
                        [cls._format_cell_value(cell, column_types.get(i)) for i, cell in enumerate(row)]
                    )

            csv_paths.append(csv_path)

        wb.close()
        return csv_paths

    @classmethod
    def _get_temporal_field_types(cls, record_type: type) -> dict[str, type]:
        """Get a mapping of field name to temporal type (dt.date, dt.time, dt.datetime) from record type."""

        hints = get_type_hints(record_type)
        result = {}
        for name, hint in hints.items():
            base = cls._extract_base_type(hint)
            if base in (dt.date, dt.time, dt.datetime):
                result[name] = base
        return result

    @classmethod
    def _extract_base_type(cls, type_hint) -> type:
        """Extract the base type from Optional[X] or X | None."""

        origin = get_origin(type_hint)
        if origin is Union or isinstance(type_hint, types.UnionType):
            args = [a for a in get_args(type_hint) if a is not type(None)]
            if args:
                return args[0]
        return type_hint

    @classmethod
    def _normalize_sheet_name(cls, name: str) -> str:
        """Remove spaces and invalid filename characters from sheet name."""
        name = name.replace(" ", "")
        name = _INVALID_SHEET_NAME_RE.sub("", name)
        return name

    @classmethod
    def _format_cell_value(cls, value, field_type: type | None = None) -> str:
        """Format a cell value for CSV output based on field type.

        CSV format: date and time as compact_str, datetime as iso_str.
        Excel format: date and time as iso_int (integer), datetime as native Excel datetime.
        """

        if value is None:
            return ""

        # Format based on field type from schema
        if field_type is dt.datetime:
            if isinstance(value, dt.datetime):
                ms = value.microsecond // 1000
                return (
                    f"{value.year:04}-{value.month:02}-{value.day:02}"
                    f"T{value.hour:02}:{value.minute:02}:{value.second:02}.{ms:03}Z"
                )
            raise RuntimeError(
                f"Expected native Excel datetime for datetime field, got {type(value).__name__}: {value}"
            )

        if field_type is dt.date:
            if isinstance(value, (int, float)):
                # iso_int stored in Excel
                return f"{int(value):08d}"
            if isinstance(value, dt.datetime):
                if value.hour != 0 or value.minute != 0 or value.second != 0 or value.microsecond != 0:
                    raise RuntimeError(
                        f"Date field has non-midnight time component: {value}. "
                        f"In Excel, store dates as iso_int (yyyymmdd integer)."
                    )
                return f"{value.year:04}{value.month:02}{value.day:02}"
            if isinstance(value, dt.date):
                return f"{value.year:04}{value.month:02}{value.day:02}"
            raise RuntimeError(f"Expected iso_int or date for date field, got {type(value).__name__}: {value}")

        if field_type is dt.time:
            if isinstance(value, (int, float)):
                # iso_int stored in Excel
                return f"{int(value):09d}"
            if isinstance(value, dt.time):
                ms = value.microsecond // 1000
                return f"{value.hour:02}{value.minute:02}{value.second:02}{ms:03}"
            raise RuntimeError(f"Expected iso_int or time for time field, got {type(value).__name__}: {value}")

        # No field type info - non-temporal field
        return str(value)
